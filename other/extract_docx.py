from docx import Document
import os

doc = Document(r'd:\Code\Vbird_ESI\other\需求描述.docx')

# Extract all paragraphs
with open(r'd:\Code\Vbird_ESI\other\需求描述_extracted.txt', 'w', encoding='utf-8') as f:
    f.write("=== PARAGRAPHS ===\n")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            f.write(f"[{i}] {p.text.strip()}\n\n")
    
    f.write("\n=== TABLES ===\n")
    for ti, t in enumerate(doc.tables):
        f.write(f"\n--- Table {ti} ---\n")
        for ri, row in enumerate(t.rows):
            cells = [c.text.strip() for c in row.cells]
            f.write(f"  Row {ri}: {cells}\n")

# Extract images
out_dir = r'd:\Code\Vbird_ESI\other\docx_images'
os.makedirs(out_dir, exist_ok=True)
for rel in doc.part.rels.values():
    if 'image' in rel.reltype:
        img_data = rel.target_part.blob
        img_name = os.path.basename(rel.target_ref)
        with open(os.path.join(out_dir, img_name), 'wb') as f:
            f.write(img_data)
        print(f"Saved: {img_name}")

print("Done!")
